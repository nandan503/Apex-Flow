#!/usr/bin/env python3
"""Build a native, editable PowerPoint deck from presentation/index.html.

    python3 presentation/build_pptx.py            # writes presentation/Apex-Flow-deck.pptx

The HTML deck is the source of truth: every slide, card, code panel, table and
speaker note in the .pptx is parsed out of it, so the two can never drift apart
in content. What cannot be carried across is click-to-reveal fragments —
python-pptx has no animation API — so every slide is fully revealed instead.

No external services, no browser: only python-pptx (pip install python-pptx).
"""
from __future__ import annotations

import html as htmllib
import math
import re
import sys
from html.parser import HTMLParser

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

SRC = "presentation/index.html"
OUT = "presentation/Apex-Flow-deck.pptx"

# ── palette, taken from presentation/deck.css (which takes it from the app) ──
NAVY = RGBColor(0x06, 0x2B, 0x57)
NAVY_DEEP = RGBColor(0x03, 0x1B, 0x39)
BLUE = RGBColor(0x07, 0x89, 0xFF)
BLUE_DIM = RGBColor(0x05, 0x6E, 0xCE)
BLUE_WASH = RGBColor(0xEE, 0xF6, 0xFF)
BG = RGBColor(0xF6, 0xF8, 0xFC)
SURFACE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x0B, 0x1F, 0x3A)
MUTED = RGBColor(0x5A, 0x6F, 0x88)
LINE = RGBColor(0xE1, 0xE8, 0xF0)
OK = RGBColor(0x08, 0xA7, 0x5A)
OK_WASH = RGBColor(0xE5, 0xF8, 0xEE)
WARN = RGBColor(0xF5, 0x9E, 0x0B)
WARN_WASH = RGBColor(0xFE, 0xF4, 0xE2)
BAD = RGBColor(0xEF, 0x33, 0x48)
BAD_WASH = RGBColor(0xFD, 0xEA, 0xEC)
INFO = RGBColor(0x69, 0x46, 0xD8)
CODE_BG = RGBColor(0x07, 0x1A, 0x30)
CODE_TXT = RGBColor(0xD7, 0xE6, 0xF7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

SANS = "Segoe UI"
MONO = "Consolas"

SLIDE_W, SLIDE_H = 13.333, 7.5
MARG = 0.5
VOID = {"br", "img", "hr", "meta", "link", "input", "path", "circle", "rect",
        "line", "use", "source", "stop", "polyline", "polygon", "ellipse"}


# ══════════════════════════════════════════════════════════════════════════
# 1 · tiny DOM
# ══════════════════════════════════════════════════════════════════════════
class Node:
    """Element node. `kids` holds elements *and* strings in document order, so a
    sentence like `a <b>b</b> c` keeps its word order (a naive "text first, then
    children" tree silently reorders every line with inline markup)."""
    __slots__ = ("tag", "attrs", "kids", "parent")

    def __init__(self, tag, attrs, parent=None):
        self.tag, self.attrs, self.parent = tag, dict(attrs), parent
        self.kids = []

    @property
    def cls(self):
        return self.attrs.get("class", "")

    def has(self, c):
        return c in self.cls.split()

    def find(self, tag=None, cls=None):
        for k in self.walk():
            if isinstance(k, str):
                continue
            if (tag is None or k.tag == tag) and (cls is None or k.has(cls)):
                return k
        return None

    def walk(self):
        for k in self.kids:
            if isinstance(k, str):
                continue
            yield k
            yield from k.walk()

    def only(self, tag=None, cls=None):
        return [k for k in self.kids if not isinstance(k, str)
                and (tag is None or k.tag == tag) and (cls is None or k.has(cls))]

    def text(self):
        out = []
        for k in self.kids:
            out.append(k if isinstance(k, str) else k.text())
        return "".join(out)


class DOM(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.root = Node("root", {})
        self.cur = self.root

    def _add(self, n):
        self.cur.kids.append(n)

    def handle_starttag(self, t, a):
        n = Node(t, a, self.cur)
        self._add(n)
        if t not in VOID:
            self.cur = n

    def handle_startendtag(self, t, a):
        self._add(Node(t, a, self.cur))

    def handle_endtag(self, t):
        c = self.cur
        while c is not self.root and c.tag != t:
            c = c.parent
        if c.parent:
            self.cur = c.parent

    def handle_data(self, d):
        if d and self.cur.kids and isinstance(self.cur.kids[-1], str):
            self.cur.kids[-1] += d
        elif d:
            self.cur.kids.append(d)


def parse(path):
    p = DOM()
    p.feed(open(path, encoding="utf-8").read())
    return [n for n in p.root.walk() if n.tag == "section" and n.has("slide")]


# ══════════════════════════════════════════════════════════════════════════
# 2 · inline markup -> runs
# ══════════════════════════════════════════════════════════════════════════
MONO_CLS = {"mono", "k2", "path", "num", "c", "s", "k", "f", "n", "p", "l"}


def runs(node, bold=False, mono=False, caps=False):
    """Flatten a node into [(text, bold, mono)] runs, in document order."""
    out = []
    for k in node.kids:
        if isinstance(k, str):
            t = k if mono else re.sub(r"[ \t]*\n[ \t]*", " ", k)
            if t:
                out.append((t, bold, mono))
            continue
        if k.tag == "br":
            out.append(("\n", bold, mono))
            continue
        if k.tag == "svg":
            continue
        b = bold or k.tag in ("b", "strong")
        m = mono or k.has("mono") or (k.tag == "span" and bool(set(k.cls.split()) & MONO_CLS))
        out.extend(runs(k, b, m, caps))
    if caps:
        out = [(t.upper(), b, m) for t, b, m in out]
    return [(t, b, m) for t, b, m in out if t]


CHIP_SEP = "   ·   "
OVERFLOWS = []          # estimated card overflows, reported by main()
FONTS = []              # (slide_title, card, chosen pt) for every card rendered


def chip_text(node):
    """`div.chips` / `div.row` hold sibling pills with no whitespace between
    them in the source (CSS supplies the gap), so join them explicitly."""
    parts = [text(k) for k in elems(node) if text(k)]
    return CHIP_SEP.join(parts) if len(parts) > 1 else (parts[0] if parts else text(node))


def elems(node):
    """Element children only (the DOM stores interleaved text as str kids)."""
    return [k for k in node.kids if not isinstance(k, str)]


def text(node, sep=" "):
    if node is None:
        return ""
    return re.sub(r"\s+", sep, node.text()).strip()


# ══════════════════════════════════════════════════════════════════════════
# 3 · geometry helpers
# ══════════════════════════════════════════════════════════════════════════
def cpl(width_in, fs):
    """Rough characters per line for a given box width and font size."""
    return max(6, (width_in * 72.0) / (fs * 0.492))


def block_h(txt, width_in, fs, lh=1.30, gap=3.0):
    lines = 0
    for para in txt.split("\n"):
        lines += max(1, math.ceil(len(para) / cpl(width_in, fs)))
    return lines * fs * lh + gap


# ══════════════════════════════════════════════════════════════════════════
# 4 · shape painters
# ══════════════════════════════════════════════════════════════════════════
def box(slide, x, y, w, h, fill=None, line=None, lw=0.75, radius=0.055, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    sp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            sp.adjustments[0] = radius
        except Exception:
            pass
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = Pt(lw)
    sp.shadow.inherit = False
    sp.text_frame.word_wrap = True
    for m in ("margin_left", "margin_right"):
        setattr(sp.text_frame, m, Inches(0.09))
    sp.text_frame.margin_top = Inches(0.06)
    sp.text_frame.margin_bottom = Inches(0.05)
    return sp


def put(tf, items, fs, color=INK, first=False, align=PP_ALIGN.LEFT, space_before=0.0,
        line_sp=1.0, mono_color=None, bullet=False):
    """Write runs into a text frame. items = list of run-lists (one per paragraph)."""
    for i, rr in enumerate(items):
        p = tf.paragraphs[0] if (first and i == 0 and not tf.paragraphs[0].runs) else tf.add_paragraph()
        p.alignment = align
        p.space_before = Pt(space_before if i or not first else 0)
        p.space_after = Pt(0)
        p.line_spacing = line_sp
        for t, b, m in rr:
            for j, chunk in enumerate(t.split("\n")):
                if j:
                    p = tf.add_paragraph()
                    p.alignment = align
                    p.space_before = Pt(0)
                    p.space_after = Pt(0)
                    p.line_spacing = line_sp
                if not chunk:
                    continue
                r = p.add_run()
                r.text = chunk
                r.font.size = Pt(fs)
                r.font.name = MONO if m else SANS
                r.font.bold = b
                r.font.color.rgb = (mono_color or BLUE_DIM) if m else color
        if bullet:
            p.text = p.text  # keep simple; bullet glyph is prepended by caller
    return tf


# ══════════════════════════════════════════════════════════════════════════
# 5 · block renderers
# ══════════════════════════════════════════════════════════════════════════
VARIANT = {"ok": (OK, OK_WASH), "warn": (WARN, WARN_WASH), "no": (BAD, BAD_WASH),
           "part": (WARN, WARN_WASH), "info": (INFO, RGBColor(0xEF, 0xEB, 0xFD))}


def card_variant(card):
    for k in ("ok", "warn", "no", "info"):
        if card.has(k):
            return VARIANT[k]
    return None


def render_code(slide, code, x, y, w, avail_h, fs=None):
    """Dark code panel: header strip + <pre> lines. Returns height used."""
    hdr = code.find("header")
    path = text(hdr.find(cls="path")) if hdr else ""
    note = text(hdr.find(cls="note")) if hdr else ""
    pre = code.find("pre")
    lines = [text(l) for l in pre.only(cls="l")] if pre else []
    if pre is not None and not lines:
        lines = text(pre, "\n").split("\n")
    if fs is None:
        # shrink until the longest line and the line count both fit
        for cand in (9.5, 9.0, 8.5, 8.0, 7.5, 7.0):
            longest = max((len(l) for l in lines), default=0)
            if longest * cand * 0.492 / 72.0 <= w - 0.24 and \
               len(lines) * cand * 1.42 + 0.30 <= avail_h:
                fs = cand
                break
        else:
            fs = 6.5
    head_h = 0.24 if (path or note) else 0.06
    body_h = max(len(lines) * fs * 1.42 / 72.0, 0.12)
    h = min(avail_h, head_h + body_h + 0.14)
    panel = box(slide, x, y, w, h, fill=CODE_BG, line=RGBColor(0x1B, 0x3A, 0x5C), lw=0.5)
    tf = panel.text_frame
    tf.margin_top = Inches(0.05)
    if path or note:
        p = tf.paragraphs[0]
        p.space_after = Pt(2)
        r = p.add_run(); r.text = ("● " if path else "") + path
        r.font.size = Pt(fs - 0.5); r.font.name = MONO; r.font.color.rgb = RGBColor(0x7F, 0xA8, 0xD0)
        if note:
            r2 = p.add_run(); r2.text = "   " + note
            r2.font.size = Pt(fs - 1.0); r2.font.name = SANS; r2.font.color.rgb = RGBColor(0x9F, 0xB6, 0xCE)
        first = False
    else:
        first = True
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if (first and i == 0) else tf.add_paragraph()
        p.space_before = Pt(0); p.space_after = Pt(0); p.line_spacing = 1.0
        r = p.add_run(); r.text = ln if ln else " "
        r.font.size = Pt(fs); r.font.name = MONO
        r.font.color.rgb = CODE_TXT
        first = False
    return h


def render_table(slide, tbl, x, y, w, avail_h):
    rows = [r for r in tbl.walk() if r.tag == "tr"]
    if not rows:
        return 0.0
    ncol = max(len([c for c in elems(r) if c.tag in ("td", "th")]) for r in rows)
    nrow = len(rows)
    h = min(avail_h, max(0.5, nrow * 0.245))
    gf = slide.shapes.add_table(nrow, ncol, Inches(x), Inches(y), Inches(w), Inches(h))
    t = gf.table
    t.first_row = True
    widths = [1.0] * ncol
    for ri, r in enumerate(rows):
        cells = [c for c in elems(r) if c.tag in ("td", "th")]
        for ci in range(ncol):
            cell = t.cell(ri, ci)
            cell.margin_left = Inches(0.06); cell.margin_right = Inches(0.05)
            cell.margin_top = Inches(0.02); cell.margin_bottom = Inches(0.02)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame; tf.word_wrap = True
            src = cells[ci] if ci < len(cells) else None
            rr = runs(src) if src is not None else [("(—)", False, False)]
            p = tf.paragraphs[0]
            p.line_spacing = 1.0
            for txt_, b, m in rr:
                run = p.add_run(); run.text = txt_
                run.font.size = Pt(8.6 if ri else 8.8)
                run.font.name = MONO if m else SANS
                run.font.bold = b or ri == 0
                run.font.color.rgb = WHITE if ri == 0 else (BLUE_DIM if m else INK)
            if src is not None and src.has("c"):
                p.alignment = PP_ALIGN.CENTER
            cell.fill.solid()
            cell.fill.fore_color.rgb = NAVY if ri == 0 else (SURFACE if ri % 2 else RGBColor(0xF2, 0xF6, 0xFB))
    return h


def render_chain(slide, ch, x, y, w, fs=8.2):
    items = [text(s) for s in ch.only("span") if not s.has("ic")]
    if not items:
        return 0.0
    label = "  →  ".join(items)
    h = block_h(label, w, fs, lh=1.25, gap=0) + 0.14
    b = box(slide, x, y, w, h, fill=BLUE_WASH, line=RGBColor(0xC9, 0xE1, 0xFF), lw=0.5)
    put(b.text_frame, [[(label, False, True)]], fs, color=NAVY, first=True, line_sp=1.15)
    return h


# ══════════════════════════════════════════════════════════════════════════
# 6 · card renderer with font auto-shrink
# ══════════════════════════════════════════════════════════════════════════
def card_blocks(card):
    """Ordered content blocks of a card: (kind, node)."""
    out = []
    for k in elems(card):
        if k.tag == "h3":
            out.append(("h", k))
        elif k.tag in ("p",):
            out.append(("p", k))
        elif k.tag in ("ul", "ol"):
            out.append(("list", k))
        elif k.tag == "div":
            if k.has("code"):
                out.append(("code", k))
            elif k.has("chain"):
                out.append(("chain", k))
            elif k.has("mono"):
                out.append(("mono", k))
            elif k.has("stat"):
                out.append(("stat", k))
            elif k.has("small"):
                out.append(("small", k))
            elif k.has("chips"):
                out.append(("chips", k))
            elif k.has("row"):
                out.append(("row", k))
            else:
                out.append(("p", k))
        elif k.tag == "table":
            out.append(("table", k))
    return out


def card_height(card, w, fs):
    """Estimated content height of a card body, in INCHES (block_h returns points;
    mixing the two units here made every card look oversized and shrank all type
    to the floor — keep this function in inches end to end)."""
    total = 0.16
    for kind, n in card_blocks(card):
        if kind == "h":
            total += block_h(text(n), w - 0.26, fs + 1.0, lh=1.10, gap=0) / 72.0 + 0.038
        elif kind == "list":
            for li in n.only("li"):
                total += block_h(text(li), w - 0.34, fs, lh=1.18, gap=0) / 72.0 + 0.017
        elif kind == "code":
            pre = n.find("pre")
            lines = [text(l) for l in pre.only(cls="l")] if pre is not None else []
            total += (0.24 if n.find("header") is not None else 0.06) \
                + max(len(lines), 1) * fs * 1.42 / 72.0 + 0.12
        elif kind == "chain":
            t = "   →   ".join(text(s) for s in n.only("span") if not s.has("ic"))
            total += block_h(t, w - 0.26, fs - 1.5, lh=1.15, gap=0) / 72.0 + 0.05
        elif kind == "table":
            total += max(0.4, len([r for r in n.walk() if r.tag == "tr"]) * 0.245)
        elif kind in ("chips", "row"):
            total += block_h(chip_text(n), w - 0.26, fs - 0.8, lh=1.20, gap=0) / 72.0 + 0.05
        else:
            size = fs - (1.2 if kind in ("small", "stat") else 0)
            total += block_h(text(n), w - 0.26, size, lh=1.18, gap=0) / 72.0 + 0.042
    return total


def pick_font(card, w, avail_h, base=10.8, floor=7.4):
    """Largest size from a fixed ladder that still fits — a ladder rather than a
    float decrement, so the result never drifts a step past the floor."""
    for cand in [round(base - 0.2 * i, 1) for i in range(int((base - floor) / 0.2) + 1)]:
        if card_height(card, w, cand) <= avail_h:
            return cand
    return floor


def _style(kind, fs, var):
    if kind == "h":
        return fs + 1.0, (var[0] if var else NAVY), 1.05
    if kind in ("small", "stat"):
        return fs - 1.2, MUTED, 1.12
    if kind == "mono":
        return fs - 0.4, BLUE_DIM, 1.20
    if kind == "chain":
        return fs - 1.5, NAVY, 1.15
    if kind in ("chips", "row"):
        return fs - 0.8, BLUE_DIM, 1.20
    return fs, INK, 1.16


def render_card(slide, card, x, y, w, h):
    """Card = background rect + absolutely positioned blocks. Positioning (rather
    than one auto-flowing text frame) is what lets a code panel inside a card stay
    a real dark panel instead of collapsing into jammed prose."""
    var = card_variant(card)
    fill = var[1] if var else SURFACE
    edge = RGBColor(0xD5, 0xE2, 0xEF) if not var else var[0]
    sh = box(slide, x, y, w, h, fill=fill, line=edge, lw=0.75)
    sh.text_frame.text = ""
    if var:
        bar = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x + 0.012),
                                     Inches(y + 0.09), Inches(0.045), Inches(max(0.1, h - 0.18)))
        bar.adjustments[0] = 0.5
        bar.fill.solid(); bar.fill.fore_color.rgb = var[0]
        bar.line.fill.background(); bar.shadow.inherit = False

    fs = pick_font(card, w, h)
    x0, w0 = x + 0.14, w - 0.26
    cy, limit = y + 0.10, y + h - 0.08
    overflow = 0.0
    for kind, n in card_blocks(card):
        # No early exit: dropping a block to keep a card tidy would silently lose
        # content. If the estimate still overflows, the block is clamped and the
        # shortfall is reported by main() so it can be eyeballed, not hidden.
        room = max(0.16, limit - cy)
        if kind == "code":
            cy += render_code(slide, n, x0, cy, w0, room) + 0.07
            continue
        if kind == "table":
            cy += render_table(slide, n, x0, cy, w0, room) + 0.07
            continue
        size, color, lsp = _style(kind, fs, var)
        if kind == "list":
            items, plain = [], []
            ordered = n.tag == "ol"
            for i, li in enumerate(n.only("li"), 1):
                mark = f"{i}.  " if ordered else "▸  "
                items.append([(mark, False, True)] + runs(li))
                plain.append(mark + text(li))
            body = "\n".join(plain)
        elif kind in ("chips", "row"):
            t = chip_text(n)
            items, body = [[(t, False, True)]], t
        elif kind == "chain":
            t = "   →   ".join(text(s) for s in n.only("span") if not s.has("ic"))
            items, body = [[(t, False, True)]], t
        elif kind == "table":
            continue
        else:
            rr = runs(n)
            if not any(t.strip() for t, _, _ in rr):
                continue
            items, body = [rr], text(n)
        hh = block_h(body, w0, size, lh=lsp, gap=0) / 72.0 + 0.035
        if hh > room:
            overflow += hh - room
            hh = room
        b = box(slide, x0, cy, w0, hh, fill=None, line=None)
        tf = b.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        put(tf, items, size, color=color, first=True, line_sp=lsp)
        cy += hh + (0.040 if kind in ("h",) else 0.058)
    FONTS.append((text(card.find("h3"))[:34], fs))
    if overflow > 0.05:
        OVERFLOWS.append(f"{card.attrs.get('data-title','')}"
                         f"{text(card.find('h3'))[:38]!r} over by {overflow:.2f} in at {fs} pt")
    return fs


# ══════════════════════════════════════════════════════════════════════════
# 7 · flow diagrams (node chains)
# ══════════════════════════════════════════════════════════════════════════
def grid_cols(node):
    """Read grid-template-columns. CSS separates tracks with whitespace, and
    repeat(N, X) contains a comma, so handle that before splitting."""
    st = node.attrs.get("style", "")
    m = re.search(r"grid-template-columns:\s*([^;\"]+)", st)
    if not m:
        return None
    spec = m.group(1).strip()
    rep = re.match(r"repeat\(\s*(\d+)\s*,\s*([^)]*)\)", spec)
    if rep:
        inner = rep.group(2).strip()
        base = float(inner.rstrip("fr")) if inner.rstrip("fr") else 1.0
        return [base] * int(rep.group(1))
    tracks = []
    for tok in spec.split():
        t = tok.rstrip(",").strip()
        if not t:
            continue
        num = t[:-2] if t.endswith("fr") else t
        try:
            tracks.append(float(num))
        except ValueError:
            tracks.append(1.0)          # auto / min-content / px — treat as equal
    return tracks or [1.0]


def render_flow(slide, flow, x, y, w, h):
    """Sequence of .node boxes with arrows between them, wrapping into rows."""
    nodes = [n for n in flow.walk() if n.tag == "div" and n.has("node")]
    if not nodes:
        return 0.0
    per_row = min(len(nodes), 5)
    rows = [nodes[i:i + per_row] for i in range(0, len(nodes), per_row)]
    gap = 0.26
    row_h = (h - gap * (len(rows) - 1)) / len(rows)
    for ri, row in enumerate(rows):
        n = len(row)
        total_gap = gap * (n - 1)
        cw = (w - total_gap) / n
        for ci, nd in enumerate(row):
            cx = x + ci * (cw + gap)
            cy = y + ri * (row_h + gap)
            b = box(slide, cx, cy, cw, row_h, fill=SURFACE, line=RGBColor(0xD5, 0xE2, 0xEF), lw=0.75)
            tf = b.text_frame
            tf.margin_left = Inches(0.1); tf.margin_right = Inches(0.08)
            t = nd.find(cls="t"); m = nd.find(cls="m"); d = nd.find(cls="d")
            first = True
            if t is not None:
                put(tf, [runs(t)], 10.2, color=NAVY, first=True, line_sp=1.0)
                first = False
            if m is not None:
                put(tf, [runs(m, mono=True)], 7.8, first=first, space_before=1.5, line_sp=1.05)
                first = False
            if d is not None:
                put(tf, [runs(d)], 8.4, color=MUTED, first=first, space_before=2.5, line_sp=1.12)
                first = False
            chips = [text(s) for s in nd.only("span") if s.has("chip") and text(s)]
            if chips:
                put(tf, [[(CHIP_SEP.join(chips), False, True)]], 7.6,
                    first=first, space_before=2.5, line_sp=1.10)
            if ci < n - 1:
                ar = slide.shapes.add_textbox(Inches(cx + cw + 0.01), Inches(cy + row_h / 2 - 0.12),
                                              Inches(gap - 0.02), Inches(0.24))
                ar.shadow.inherit = False
                atf = ar.text_frame
                atf.margin_left = atf.margin_right = atf.margin_top = atf.margin_bottom = 0
                p = atf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
                r = p.add_run(); r.text = "→"
                r.font.size = Pt(13); r.font.bold = True; r.font.color.rgb = BLUE
    return h


# ══════════════════════════════════════════════════════════════════════════
# 8 · slide assembly
# ══════════════════════════════════════════════════════════════════════════
def _est_h(node, w, fs=10.6):
    """Estimated natural height (inches) of a child, used to divide space."""
    if node.tag == "table":
        return max(0.4, len([r for r in node.walk() if r.tag == "tr"]) * 0.245)
    if node.has("card"):
        return card_height(node, w, fs) + 0.14
    if node.has("code"):
        pre = node.find("pre")
        n = len(pre.only(cls="l")) if pre is not None else 3
        return (0.24 if node.find("header") else 0.06) + n * fs * 1.42 / 72.0 + 0.14
    if node.has("flow") or any(k.has("node") for k in elems(node)):
        return 1.5
    return max(0.25, block_h(text(node), w, fs, lh=1.3, gap=0) / 72.0 + 0.1)


def layout_children(container, x, y, w, h, slide, gap=0.16):
    """Place the direct children of a grid/stack/flow into the given rect. Space is
    divided by *estimated content height*, not by equal slices, so a dense card is
    not forced to shrink its type below the deck's legibility floor."""
    kids = [k for k in elems(container)
            if k.tag in ("div", "table", "p", "h1", "h2", "h3") and not k.has("evidence")
            and not k.has("band")]
    if not kids:
        return
    if container.has("flow") or any(k.has("node") for k in kids):
        render_flow(slide, container, x, y, w, h)
        return
    if container.has("stack"):
        weights = [_est_h(k, w) for k in kids]
        tot = sum(weights) or 1.0
        avail = h - gap * (len(kids) - 1)
        cy = y
        for k, wt in zip(kids, weights):
            ch = max(0.3, avail * wt / tot)
            place(k, x, cy, w, ch, slide)
            cy += ch + gap
        return
    cols = grid_cols(container) or [1.0] * (2 if len(kids) > 1 else 1)
    ncol = len(cols)
    nrow = math.ceil(len(kids) / ncol)
    tw = sum(cols)
    colw = [w * c / tw for c in cols]
    # row heights weighted by the tallest card in each row
    roww = []
    for r in range(nrow):
        row = kids[r * ncol:(r + 1) * ncol]
        roww.append(max([_est_h(k, colw[min(i, len(colw) - 1)]) for i, k in enumerate(row)] or [1.0]))
    tot = sum(roww) or 1.0
    avail = h - gap * (nrow - 1)
    ys, cy = [], y
    for r in range(nrow):
        rh = max(0.35, avail * roww[r] / tot)
        ys.append((cy, rh))
        cy += rh + gap
    for i, k in enumerate(kids):
        r, cc = divmod(i, ncol)
        cx = x + sum(colw[:cc]) + gap * cc
        cw = colw[cc] if cc < len(colw) else colw[-1]
        ry, rh = ys[r]
        place(k, cx, ry, cw, rh, slide)


BLOCK_TAGS = ("div", "table", "ul", "ol", "p", "section")
LEAF_CLS = ("stat", "small", "mono", "chips", "row", "meta", "chip", "tag", "num")


def is_leaf(node):
    """A div whose children are inline-only (or which we always render as text)."""
    if node.tag != "div":
        return node.tag in ("p", "h1", "h2", "h3", "h4")
    if set(node.cls.split()) & set(LEAF_CLS):
        return True
    return not any(k.tag in BLOCK_TAGS for k in elems(node))


def place(node, x, y, w, h, slide):
    if node.tag == "table":
        render_table(slide, node, x, y, w, h)
        return
    if node.has("card"):
        render_card(slide, node, x, y, w, h)
        return
    if node.has("code"):
        render_code(slide, node, x, y, w, h)
        return
    if node.has("node"):
        render_flow(slide, node, x, y, w, h)
        return
    if is_leaf(node):
        txt = chip_text(node) if (node.has("chips") or node.has("row")) else text(node)
        if not txt.strip():
            return
        b = box(slide, x, y, w, h, fill=None, line=None)
        tf = b.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        mono = node.has("mono") or node.has("chip")
        small = node.has("small") or node.has("stat")
        put(tf, [[(txt, node.has("chip") is False and node.tag in ("h1", "h2", "h3"), mono)]],
            9.6 if node.has("chip") or node.has("row") or node.has("chips") else (9.8 if small else 11.2),
            color=MUTED if small else (BLUE_DIM if mono else INK), first=True,
            line_sp=1.15 if not small else 1.12)
        return
    if node.has("flow") or any(k.has("node") for k in elems(node)):
        render_flow(slide, node, x, y, w, h)
        return
    layout_children(node, x, y, w, h, slide)


def render_evidence(slide, ev, dark=False):
    """Evidence footer, walked in document order: `k2` refs are monospace, a
    `small` label that follows a number is glued to it, everything else is
    separated by a middot."""
    tb = box(slide, MARG, SLIDE_H - 0.40, SLIDE_W - 2 * MARG, 0.30, fill=None, line=None)
    tf = tb.text_frame
    tf.margin_left = Inches(0); tf.margin_top = Inches(0.02)
    p = tf.paragraphs[0]; p.line_spacing = 1.0
    prev = None
    for k in ev.kids:
        if isinstance(k, str):
            t = re.sub(r"\s+", " ", k).strip()
            if not t:
                continue
            r = p.add_run(); r.text = t
            r.font.size = Pt(8.3); r.font.name = SANS
            r.font.color.rgb = RGBColor(0x9B, 0xB6, 0xD2) if dark else MUTED
            prev = "txt"; continue
        if k.tag != "span":
            continue
        t = text(k)
        if not t:
            continue
        kind = "k2" if k.has("k2") else ("small" if k.has("small") else ("tag" if k.has("tag") else "txt"))
        glue = (kind == "small" and prev == "k2")
        sep = "" if glue else ("" if prev is None else "   ·   ")
        r = p.add_run(); r.text = sep + t
        r.font.size = Pt(8.3 if kind in ("k2", "small", "txt") else 8.6)
        r.font.name = MONO if kind == "k2" else SANS
        r.font.bold = kind == "tag"
        if kind == "tag":
            r.font.color.rgb = BAD if k.has("no") else (OK if k.has("ok") else
                               (WARN if (k.has("part") or k.has("warn")) else BLUE_DIM))
        elif kind == "k2":
            r.font.color.rgb = RGBColor(0x7F, 0xA8, 0xD0) if dark else BLUE_DIM
        else:
            r.font.color.rgb = RGBColor(0x9B, 0xB6, 0xD2) if dark else MUTED
        prev = kind


def render_title_slide(prs, sec, idx):
    """Slide 1 is not a grid of cards: wordmark, tagline, chips, identity block."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = NAVY_DEEP

    eb = sec.find("div", "eyebrow")
    if eb is not None:
        t = box(slide, MARG, 0.62, SLIDE_W - 2 * MARG, 0.26, fill=None, line=None)
        put(t.text_frame, [runs(eb, caps=True)], 10.0, color=RGBColor(0x8F, 0xC4, 0xF5), first=True)

    wm = sec.find("h1")
    y = 1.02
    if wm is not None:
        accent = wm.find("span")
        acc_txt = text(accent) if accent is not None else ""
        b = box(slide, MARG, y, SLIDE_W - 2 * MARG, 1.32, fill=None, line=None)
        p0 = b.text_frame.paragraphs[0]
        p0.line_spacing = 0.95
        for t_, bo, mo in runs(wm):
            r = p0.add_run(); r.text = t_
            r.font.size = Pt(64); r.font.bold = True; r.font.name = SANS
            r.font.color.rgb = BLUE if (acc_txt and t_.strip() == acc_txt) else WHITE
        y += 1.40

    tag = sec.find("p", "tagline") or sec.find("p", "lede")
    if tag is not None:
        h = block_h(text(tag), 10.6, 15.0, lh=1.30, gap=0) / 72.0 + 0.12
        b = box(slide, MARG, y, 10.6, h, fill=None, line=None)
        put(b.text_frame, [runs(tag)], 15.0, color=RGBColor(0xC5, 0xD8, 0xEE), first=True, line_sp=1.24)
        y += h + 0.20

    chips = sec.find("div", "row") or sec.find("div", "chips")
    if chips is not None:
        x = MARG
        for s in chips.only("span"):
            nm = text(s)
            if not nm:
                continue
            wch = 0.32 + len(nm) * 0.082
            ch = box(slide, x, y, wch, 0.30, fill=RGBColor(0x0E, 0x2C, 0x4E),
                     line=RGBColor(0x25, 0x4A, 0x72), lw=0.75, radius=0.42)
            tf = ch.text_frame
            tf.margin_top = Inches(0.02); tf.margin_bottom = Inches(0.02)
            pp = tf.paragraphs[0]; pp.alignment = PP_ALIGN.CENTER
            r = pp.add_run(); r.text = nm
            r.font.size = Pt(9.5); r.font.name = SANS; r.font.color.rgb = RGBColor(0xBF, 0xD9, 0xF2)
            x += wch + 0.11
        y += 0.54

    meta = sec.find("div", "meta")
    pending = False
    if meta is not None:
        cols = elems(meta)
        n = max(1, len(cols))
        cw = (SLIDE_W - 2 * MARG) / n
        for i, col in enumerate(cols):
            m1, m2 = col.find(cls="m1"), col.find(cls="m2")
            b = box(slide, MARG + i * cw, y, cw - 0.24, 0.62, fill=None, line=None)
            tf = b.text_frame
            tf.margin_left = tf.margin_top = 0
            put(tf, [runs(m1, caps=True) if m1 is not None else [("—", False, False)]], 8.6,
                color=RGBColor(0x7F, 0xA6, 0xCC), first=True, line_sp=1.0)
            if m2 is not None and "<" in text(m2):
                pending = True
            put(tf, [runs(m2) if m2 is not None else [("—", False, False)]], 11.5,
                color=WARN if (m2 is not None and "<" in text(m2)) else WHITE,
                space_before=2.5, line_sp=1.0)
        if pending:
            w = box(slide, MARG, y + 0.66, 10.2, 0.26, fill=None, line=None)
            put(w.text_frame, [[("⚠ placeholders still unfilled — set META at presentation/deck.js:10 before presenting", True, False)]],
                9.5, color=WARN, first=True)

    ev = sec.find("div", "evidence")
    if ev is not None:
        render_evidence(slide, ev, dark=True)
    pn = box(slide, SLIDE_W - 1.15, SLIDE_H - 0.40, 0.75, 0.26, fill=None, line=None)
    pp = pn.text_frame.paragraphs[0]; pp.alignment = PP_ALIGN.RIGHT
    r = pp.add_run(); r.text = f"{idx:02d} / 18"
    r.font.size = Pt(8.5); r.font.name = MONO; r.font.color.rgb = RGBColor(0x74, 0x93, 0xB4)
    notes = sec.attrs.get("data-notes", "")
    if notes:
        slide.notes_slide.notes_text_frame.text = re.sub(
            r"\s+", " ", htmllib.unescape(re.sub(r"<[^>]+>", "", notes))).strip()
    return slide


def build_slide(prs, sec, idx):
    if sec.attrs.get("data-title") == "Title":
        return render_title_slide(prs, sec, idx)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    dark = sec.has("dark")
    bg = slide.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = NAVY_DEEP if dark else BG

    title = sec.attrs.get("data-title", f"Slide {idx}")
    body = sec
    lede = sec.find("p", "lede")
    ev = sec.find("div", "evidence")

    # content blocks in document order, minus chrome
    blocks = [k for k in elems(body)
              if not (k.tag == "div" and (k.has("eyebrow") or k.has("evidence") or k.has("band")))
              and k.tag != "h2" and not (k.tag == "p" and k.has("lede"))]

    y = 0.30
    if dark and sec.has("active") or title == "Title":
        pass

    # eyebrow
    eb = body.find("div", "eyebrow")
    if eb is not None:
        t = box(slide, MARG, y, SLIDE_W - 2 * MARG, 0.24, fill=None, line=None)
        put(t.text_frame, [runs(eb, caps=True)], 9.5,
            color=RGBColor(0x8F, 0xC4, 0xF5) if dark else BLUE_DIM, first=True)
        y += 0.26

    # title
    h2 = body.find("h2")
    if h2 is not None:
        tt = text(h2)
        fs = 25.0 if len(tt) < 62 else (22.5 if len(tt) < 92 else 20.0)
        tb = box(slide, MARG, y, SLIDE_W - 2 * MARG, 0.55 if fs > 22 else 0.68, fill=None, line=None)
        tb.text_frame.vertical_anchor = MSO_ANCHOR.TOP
        put(tb.text_frame, [runs(h2)], fs, color=WHITE if dark else NAVY, first=True, line_sp=0.98)
        y += (0.58 if fs > 22 else 0.72)

    # lede
    if lede is not None:
        lh = block_h(text(lede), SLIDE_W - 2 * MARG, 12.5, lh=1.28, gap=0) / 72.0 + 0.10
        lb = box(slide, MARG, y, SLIDE_W - 2 * MARG, lh, fill=None, line=None)
        put(lb.text_frame, [runs(lede)], 12.5,
            color=RGBColor(0xC5, 0xD8, 0xEE) if dark else RGBColor(0x33, 0x4C, 0x6B), first=True, line_sp=1.22)
        y += lh + 0.06

    # evidence footer
    y1 = SLIDE_H - (0.46 if ev is not None else 0.26)
    if blocks:
        layout_children(blocks[0] if len(blocks) == 1 else _wrap(blocks), MARG, y,
                        SLIDE_W - 2 * MARG, max(0.6, y1 - y), slide)

    if ev is not None:
        render_evidence(slide, ev, dark=dark)

    # page number
    pn = box(slide, SLIDE_W - 1.15, SLIDE_H - 0.40, 0.75, 0.26, fill=None, line=None)
    p = pn.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.RIGHT
    r = p.add_run(); r.text = f"{idx:02d} / 18"
    r.font.size = Pt(8.5); r.font.name = MONO
    r.font.color.rgb = RGBColor(0x7E, 0x99, 0xB6) if not dark else RGBColor(0x74, 0x93, 0xB4)

    # speaker notes straight from data-notes
    notes = sec.attrs.get("data-notes", "")
    if notes:
        slide.notes_slide.notes_text_frame.text = re.sub(
            r"\s+", " ", htmllib.unescape(re.sub(r"<[^>]+>", "", notes))).strip()
    return slide


class _Wrap(Node):
    def __init__(self, kids):
        super().__init__("div", {"class": "grid grow",
                                 "style": "grid-template-columns:" + ",".join(["1fr"] * len(kids))})
        self.kids = kids


def _wrap(blocks):
    w = _Wrap(list(blocks))
    w.attrs["style"] = "grid-template-columns:1fr"
    return w


def main():
    slides = parse(SRC)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(SLIDE_W), Inches(SLIDE_H)
    for i, sec in enumerate(slides, 1):
        build_slide(prs, sec, i)
    prs.core_properties.title = "Apex Flow — technical defence"
    prs.core_properties.author = "nandan503"
    prs.core_properties.comments = ("Generated from presentation/index.html by presentation/build_pptx.py. "
                                    "All claims trace to files in this repository; see docs/PRESENTATION-NOTES.md §8.")
    prs.save(OUT)
    print(f"wrote {OUT}: {len(prs.slides._sldIdLst)} slides")
    if OVERFLOWS:
        print(f"\n⚠ {len(OVERFLOWS)} card(s) estimated to overflow their box (content kept, not dropped):")
        for o in OVERFLOWS:
            print("   ·", o)
    else:
        print("no card overflow estimated")
    if FONTS:
        sizes = sorted(f for _, f in FONTS)
        print(f"card type sizes: min {sizes[0]} pt, median {sizes[len(sizes)//2]} pt, max {sizes[-1]} pt "
              f"({sum(1 for s in sizes if s < 8)} of {len(sizes)} cards below 8 pt)")
        for name, f in FONTS:
            if f < 8:
                print(f"   · small: {f} pt — {name!r}")


if __name__ == "__main__":
    sys.exit(main())
